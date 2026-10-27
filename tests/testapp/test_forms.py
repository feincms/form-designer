import io
from datetime import datetime
from unittest import mock

import openpyxl
from django import forms
from django.contrib.auth.models import User
from django.core import mail
from django.test import TestCase, override_settings
from django.utils.text import slugify
from feincms.module.page.models import Page

from form_designer.models import FIELD_TYPES, Form, FormField, FormSubmission


def validate_honeypot(form, data, **kwargs):
    if data.get("honeypot"):
        raise forms.ValidationError("Hello honeypot")


Form.CONFIG_OPTIONS.append(
    ("honeypot", {"title": "Honeypot", "validate": validate_honeypot})
)


class FormsTest(TestCase):
    maxDiff = None

    def test_forms(self):
        form = Form.objects.create(
            title="Test contact form",
            config={"email": {"email": "info@example.com"}, "save_fs": {}},
        )
        form.fields.create(ordering=0, title="Subject", name="subject", type="text")
        form.fields.create(ordering=1, title="Email", name="email", type="email")
        form.fields.create(ordering=2, title="Body", name="body", type="longtext")
        form.fields.create(
            ordering=3,
            title="Please call me",
            name="please-call-me",
            type="checkbox",
            is_required=False,
        )
        form.fields.create(
            ordering=4,
            title="Radio Select",
            name="radio",
            type="radio",
            choices="one,two what,three",
            default_value="two what",
        )
        form.fields.create(
            ordering=5,
            title="Date",
            name="date",
            type="date",
            is_required=False,
        )
        form.fields.create(
            ordering=6,
            title="Multiple Choice",
            name="multiple-choice",
            type="multiple-select",
            choices="Choice A,Choice B,Choice C",
            is_required=False,
        )

        form_class = form.form_class()
        form_instance = form_class()

        self.assertListEqual(
            [field.name for field in form_instance],
            [
                "subject",
                "email",
                "body",
                "please-call-me",
                "radio",
                "date",
                "multiple-choice",
            ],
        )

        page = Page.objects.create(override_url="/", title="")
        page.formcontent_set.create(
            region="main",
            ordering=0,
            form=form,
            success_message="Thanks, we will get back to you",
        )

        response = self.client.get("/")

        self.assertContains(response, 'method="post"', 1)
        self.assertContains(response, f'action="#form{form.id}"', 1)
        self.assertContains(
            response,
            "<input",
            13,  # csrf, subject, email, checkbox, _formcontent, submit, radio*3, date, three multiple choice
        )
        self.assertContains(response, "<textarea", 1)
        self.assertContains(
            response, 'value="two what" required id="id_fc1-radio_1" checked', 1
        )
        self.assertContains(response, 'type="date"', 1)

        response = self.client.post("/")

        self.assertContains(response, "This field is required", 4)

        # Not this form
        response = self.client.post("/", {"_formcontent": -1})

        self.assertNotContains(response, "This field is required")

        response = self.client.post(
            "/",
            {
                "_formcontent": form.id,
                f"fc{form.id}-subject": "Test",
                f"fc{form.id}-email": "invalid",
                f"fc{form.id}-body": "Hello World",
                f"fc{form.id}-radio": "two-what",
            },
        )

        self.assertNotContains(response, "This field is required")

        self.assertContains(
            response,
            "Enter a valid email",
            1,
        )

        response = self.client.post(
            "/",
            {
                "_formcontent": form.id,
                f"fc{form.id}-subject": "Test",
                f"fc{form.id}-email": "valid@example.com",
                f"fc{form.id}-body": "Hello World",
                f"fc{form.id}-radio": "two what",
                f"fc{form.id}-date": "2022-10-02",
                f"fc{form.id}-multiple-choice": ["Choice A", "Choice C"],
            },
        )

        self.assertContains(response, "Thanks, we will get back to you", 1)

        self.assertNotContains(response, "<form")

        self.assertEqual(len(mail.outbox), 1)

        message = mail.outbox[0]

        self.assertListEqual(message.to, ["info@example.com"])
        self.assertEqual(message.from_email, "no-reply@example.com")
        self.assertEqual(message.subject, "Test contact form")
        self.assertIn("Subject:\nTest\n", message.body)
        self.assertIn("Email:\nvalid@example.com\n", message.body)
        self.assertIn("Body:\nHello World\n", message.body)
        self.assertIn("Please call me:\nØ\n", message.body)
        self.assertIn("Date:\n2022-10-02\n", message.body)

        # Exactly one submission
        submission = FormSubmission.objects.get()

        self.assertEqual(
            form.submissions_data(submissions=[submission]),
            [
                {
                    "submission": submission,
                    "data": [
                        {"name": "subject", "title": "Subject", "value": "Test"},
                        {
                            "name": "email",
                            "title": "Email",
                            "value": "valid@example.com",
                        },
                        {"name": "body", "title": "Body", "value": "Hello World"},
                        {
                            "name": "please-call-me",
                            "title": "Please call me",
                            "value": False,
                        },
                        {"name": "radio", "title": "Radio Select", "value": "two what"},
                        {"name": "date", "title": "Date", "value": "2022-10-02"},
                        {
                            "name": "multiple-choice",
                            "title": "Multiple Choice",
                            "value": ["Choice A", "Choice C"],
                        },
                    ],
                }
            ],
        )

        data = submission.formatted_data()
        self.assertIn("Multiple Choice:\nChoice A, Choice C", data)

        # Export the submission
        User.objects.create_superuser("admin", "admin@example.com", "password")
        self.client.login(username="admin", password="password")
        with mock.patch(
            "form_designer.admin.localtime",
            return_value=datetime(2026, 10, 27, 17, 59),
        ):
            response = self.client.get(
                f"/admin/form_designer/form/{submission.form_id}/export_submissions/"
            )
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            f"{slugify(submission.form.title)}_20261027_1759.xlsx",
            response["Content-Disposition"],
        )
        self.assertEqual(
            response["Content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        # Without time zone support
        with override_settings(USE_TZ=False):
            response = self.client.get(
                f"/admin/form_designer/form/{submission.form_id}/export_submissions/"
            )
        self.assertEqual(response.status_code, 200)

        response = self.client.get(
            f"/admin/form_designer/formsubmission/{submission.id}/change/"
        )
        self.assertContains(response, "<dt>Body</dt>")

        FormSubmission.objects.all().delete()
        response = self.client.get(
            f"/admin/form_designer/form/{submission.form_id}/export_submissions/"
        )
        self.assertRedirects(
            response, f"/admin/form_designer/form/{submission.form_id}/change/"
        )

    def test_admin(self):
        User.objects.create_superuser("admin", "admin@example.com", "password")
        self.client.login(username="admin", password="password")

        response = self.client.get("/admin/form_designer/form/add/")
        self.assertEqual(response.status_code, 200)  # Oh well...

        data = {
            "title": "Test form",
            "_is_active_save_fs": "on",
            "_is_active_email": "on",
            "initial-config": "{}",
            "config": "{}",
            "email_email": "bla@example.com,blabbb@example.com",
            "fields-TOTAL_FORMS": 9,
            "fields-INITIAL_FORMS": 0,
            "fields-MIN_NUM_FORMS": 0,
            "fields-MAX_NUM_FORMS": 1000,
        }
        for i in range(9):
            data[f"fields-{i}-ordering"] = (i + 1) * 10
            data[f"fields-{i}-title"] = f"title-{i}"
            data[f"fields-{i}-name"] = f"name-{i}"
            data[f"fields-{i}-type"] = FIELD_TYPES[i]["type"]
            data[f"fields-{i}-choices"] = ""

        # Validation failure because of missing choices
        response = self.client.post("/admin/form_designer/form/add/", data)
        self.assertEqual(response.status_code, 200)

        for i in range(9):
            if FIELD_TYPES[i]["type"] in {"select", "radio", "multiple-select"}:
                data[f"fields-{i}-choices"] = "a,b,c,d"

        data["fields-3-choices"] = "invalid"
        # Validation failure because of choices where there should be none
        response = self.client.post("/admin/form_designer/form/add/", data)
        self.assertEqual(response.status_code, 200)

        data["fields-3-choices"] = ""
        # Text fields accept choices, used as placeholder
        data["fields-0-choices"] = "Placeholder"
        response = self.client.post("/admin/form_designer/form/add/", data)
        self.assertRedirects(response, "/admin/form_designer/form/")

        data["fields-4-type"] = ""  # No crash, but no success either
        response = self.client.post("/admin/form_designer/form/add/", data)
        self.assertEqual(response.status_code, 200)

        # Basic smoke test
        form = Form.objects.get()
        self.assertEqual(set(form.config), {"email", "save_fs"})
        form_class = form.form_class()

        self.assertEqual(form.title, "Test form")
        self.assertEqual(len(form_class().fields), 9)
        self.assertEqual(
            form_class().fields["name-0"].widget.attrs["placeholder"], "Placeholder"
        )

        # The additional formsets exist
        response = self.client.get(f"/admin/form_designer/form/{form.id}/change/")
        self.assertContains(response, 'id="id_email_email"')
        self.assertContains(response, "blabbb@example.com")

    def test_honeypot(self):
        form = Form.objects.create(
            title="Test honeypot form",
            config={"honeypot": {}},
        )
        form.fields.create(ordering=0, title="Subject", name="subject", type="text")
        form.fields.create(
            ordering=0,
            title="honeypot",
            name="honeypot",
            type="hidden",
            is_required=False,
        )

        page = Page.objects.create(override_url="/", title="")
        page.formcontent_set.create(
            region="main",
            ordering=0,
            form=form,
            success_message="Thanks, we will get back to you",
        )

        response = self.client.post(
            "/",
            {
                "_formcontent": form.id,
                f"fc{form.id}-subject": "Test",
                f"fc{form.id}-honeypot": "honey",
            },
        )

        self.assertContains(response, "Hello honeypot")

    def test_submission(self):
        form = Form.objects.create(
            title="Test contact form",
            config={"email": {"email": "info@example.com"}, "save_fs": {}},
        )
        form.fields.create(ordering=0, title="Subject", name="subject", type="text")
        form.fields.create(
            ordering=1, title="Email", name="email", _old_name="e mail", type="email"
        )

        s1 = FormSubmission.objects.create(
            form=form,
            data={
                "subject": "blub",
                "email": "a@example.com",
            },
        )
        s2 = FormSubmission.objects.create(
            form=form,
            data={
                "subject": "blub",
                "e mail": "a@example.com",
            },
        )

        form.fields.create(ordering=2, title="test", name="test", type="text")

        self.assertEqual(
            form.submissions_data(),
            [
                {
                    "submission": s2,
                    "data": [
                        {"name": "subject", "title": "Subject", "value": "blub"},
                        {"name": "email", "title": "Email", "value": "a@example.com"},
                        {"name": "test", "title": "test", "value": None},
                        {
                            "name": "e mail",
                            "title": "e mail (removed field)",
                            "value": "a@example.com",
                        },
                    ],
                },
                {
                    "submission": s1,
                    "data": [
                        {"name": "subject", "title": "Subject", "value": "blub"},
                        {"name": "email", "title": "Email", "value": "a@example.com"},
                        {"name": "test", "title": "test", "value": None},
                        {
                            "name": "e mail",
                            "title": "e mail (removed field)",
                            "value": None,
                        },
                    ],
                },
            ],
        )

    def test_email_to_author(self):
        form = Form.objects.create(
            title="Test contact form",
            config={
                "email": {"email": "info@example.com", "author_email_field": "mmm"}
            },
        )
        form.fields.create(ordering=0, title="Subject", name="subject", type="text")
        form.fields.create(ordering=1, title="Email", name="mmm", type="email")

        page = Page.objects.create(override_url="/", title="")
        page.formcontent_set.create(
            region="main",
            ordering=0,
            form=form,
            success_message="Thanks, we will get back to you",
        )

        response = self.client.post(
            "/",
            {
                "_formcontent": form.id,
                f"fc{form.id}-subject": "Test",
                f"fc{form.id}-mmm": "test@example.org",
            },
        )
        self.assertContains(response, "Thanks, we will get back to you", 1)

        self.assertEqual(len(mail.outbox), 1)

        message = mail.outbox[0]

        self.assertListEqual(message.to, ["info@example.com"])
        self.assertListEqual(message.cc, ["test@example.org"])
        self.assertEqual(message.from_email, "no-reply@example.com")
        self.assertEqual(message.subject, "Test contact form")
        self.assertIn("Subject:\nTest\n", message.body)
        self.assertIn("Email:\ntest@example.org\n", message.body)

    def _create_email_form(self, email_config):
        form = Form.objects.create(
            title="Test contact form", config={"email": email_config}
        )
        form.fields.create(ordering=0, title="Subject", name="subject", type="text")
        form.fields.create(
            ordering=1, title="Email", name="mmm", type="email", is_required=False
        )
        page = Page.objects.create(override_url="/", title="")
        page.formcontent_set.create(
            region="main", ordering=0, form=form, success_message="Thanks"
        )
        return form

    def _post_email_form(self, form, data):
        mail.outbox.clear()
        response = self.client.post(
            "/",
            {"_formcontent": form.id}
            | {f"fc{form.id}-{key}": value for key, value in data.items()},
        )
        self.assertContains(response, "Thanks", 1)
        self.assertEqual(len(mail.outbox), 1)
        return mail.outbox[0]

    def test_email_subject(self):
        cases = [
            ("Contact: %(subject)s", {"subject": "Hello"}, "Contact: Hello"),
            ("Contact: %(subject)s %(mmm)s", {"subject": "Hi"}, "Contact: Hi"),
            ("Contact: %(nonexistent)s", {"subject": "Hi"}, "Contact:"),
            ("100% sure", {"subject": "Hi"}, "100% sure"),
            ("%(subject)d", {"subject": "Hi"}, "Test contact form"),
            ("%(subject", {"subject": "Hi"}, "Test contact form"),
            ("100%% sure", {"subject": "Hi"}, "100% sure"),
            ("%(subject)s", {"subject": "a\nb"}, "a b"),
            ("", {"subject": "Hi"}, "Test contact form"),
        ]
        form = self._create_email_form({"email": "info@example.com"})
        for template, data, expected in cases:
            with self.subTest(template=template, data=data):
                form.config["email"]["subject"] = template
                form.save()
                message = self._post_email_form(form, data)
                self.assertEqual(message.subject, expected)

    @override_settings(FORCE_MAIL_RECIPIENT="test@example.net")
    def test_email_force_recipient(self):
        form = self._create_email_form(
            {
                "email": "info@example.com, other@example.com",
                "author_email_field": "mmm",
            }
        )
        message = self._post_email_form(
            form, {"subject": "Test", "mmm": "author@example.org"}
        )
        self.assertListEqual(message.to, ["test@example.net"])
        self.assertListEqual(message.cc, [])

    @override_settings(FORCE_MAIL_RECIPIENT=["a@example.net", "b@example.net"])
    def test_email_force_recipient_list(self):
        form = self._create_email_form({"email": "info@example.com"})
        message = self._post_email_form(form, {"subject": "Test"})
        self.assertListEqual(message.to, ["a@example.net", "b@example.net"])

    def test_choices(self):
        """Legacy slugified choices should also use the pretty format"""

        form = Form.objects.create(
            title="Test contact form",
            config={"save_fs": {}},
        )
        form.fields.create(
            ordering=6,
            title="Multiple Choice",
            name="multiple-choice",
            type="multiple-select",
            choices="Choice A,Choice B,Choice C",
            is_required=False,
        )

        s1 = FormSubmission.objects.create(
            form=form,
            data={
                "multiple-choice": ["Choice A", "Choice C"],
            },
        )
        self.assertIn("Multiple Choice:\nChoice A, Choice C", s1.formatted_data())

        s2 = FormSubmission.objects.create(
            form=form,
            data={
                "multiple-choice": ["choice-a", "choice-c"],
            },
        )
        self.assertIn("Multiple Choice:\nChoice A, Choice C", s2.formatted_data())

    def test_is_choicefield(self):
        self.assertEqual(
            {
                field_type["type"]
                for field_type in FIELD_TYPES
                if FormField(type=field_type["type"]).is_choicefield()
            },
            {"select", "radio", "multiple-select", "statictext"},
        )

    def test_pipe_choices(self):
        """Choices are split on pipes if present, so they may contain commas"""

        form = Form.objects.create(title="Test pipe choices form")
        comma = form.fields.create(
            ordering=0,
            title="Comma",
            name="comma",
            type="radio",
            choices="A, B,C",
        )
        pipe = form.fields.create(
            ordering=1,
            title="Pipe",
            name="pipe",
            type="radio",
            choices="Yes, please | No, thanks|Maybe",
        )

        self.assertEqual(comma.get_choices(), (("A", "A"), ("B", "B"), ("C", "C")))
        self.assertEqual(
            pipe.get_choices(),
            (
                ("Yes, please", "Yes, please"),
                ("No, thanks", "No, thanks"),
                ("Maybe", "Maybe"),
            ),
        )
        self.assertEqual(
            form.form_class()().fields["pipe"].choices, list(pipe.get_choices())
        )

    def test_omitted_fields(self):
        form = Form.objects.create(title="Omit form", config={"save_fs": {}})
        form.fields.create(ordering=0, title="Subject", name="subject", type="text")
        form.fields.create(ordering=1, title="Secret", name="secret", type="text")
        form.fields.create(ordering=2, title="Captcha", name="captcha", type="text")
        form.fields.create(ordering=3, title="Private", name="_private", type="text")
        # Form field keys are slugified, "_private" is submitted as "private"
        form_instance = form.form_class()(
            data={
                "subject": "Hello",
                "secret": "s3cr3t",
                "captcha": "token",
                "private": "hidden",
            }
        )
        self.assertTrue(form_instance.is_valid())
        submission = FormSubmission.objects.create(
            form=form, data=form_instance.cleaned_data
        )

        data = submission.formatted_data()
        self.assertIn("Subject:\nHello\n", data)
        self.assertIn("Secret:\ns3cr3t\n", data)
        self.assertNotIn("Captcha", data)
        self.assertNotIn("Private", data)
        self.assertNotIn("private", data)

        data = submission.formatted_data(omit_fields=["secret"])
        self.assertIn("Subject:\nHello\n", data)
        self.assertNotIn("Secret", data)

        data = submission.formatted_data(omit_fields=["sub*"])
        self.assertNotIn("Subject", data)
        self.assertIn("Secret:\ns3cr3t\n", data)

        data = submission.formatted_data(html=True, omit_fields=("secret",))
        self.assertIn("<dt>Subject</dt>", data)
        self.assertNotIn("Secret", data)
        self.assertNotIn("Captcha", data)

        with override_settings(FORM_DESIGNER_OMIT_FIELDS=["secret"]):
            data = submission.formatted_data()
        self.assertIn("Captcha:\ntoken\n", data)
        self.assertIn("Private:\nhidden\n", data)
        self.assertNotIn("Secret", data)

        # Old behaviour: nothing omitted
        with override_settings(FORM_DESIGNER_OMIT_FIELDS=()):
            data = submission.formatted_data()
        self.assertIn("Captcha:\ntoken\n", data)
        self.assertIn("Private:\nhidden\n", data)
        self.assertIn("Secret:\ns3cr3t\n", data)

        User.objects.create_superuser("admin", "admin@example.com", "password")
        self.client.login(username="admin", password="password")
        response = self.client.get(
            f"/admin/form_designer/form/{form.id}/export_submissions/"
        )
        self.assertEqual(response.status_code, 200)
        sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        self.assertEqual(rows[0][:2], ["Subject", "Secret"])
        self.assertEqual(rows[1][:2], ["subject", "secret"])
        self.assertEqual(rows[2][:2], ["Hello", "s3cr3t"])
        self.assertEqual(len(rows[0]), 4)  # + submitted at, URL
        self.assertNotIn("captcha", rows[1])
        self.assertNotIn("_private", rows[1])

    def test_statictext(self):
        form = Form.objects.create(title="Static form", config={"save_fs": {}})
        form.fields.create(ordering=0, title="Subject", name="subject", type="text")
        form.fields.create(
            ordering=1,
            title="Hello <b>world</b>",
            name="intro",
            type="statictext",
            choices="lead|highlight",
        )
        form.fields.create(
            ordering=2,
            title="Paragraph",
            name="para",
            type="statictext",
            default_value="p",
            is_required=True,
        )

        FormField(type="statictext", default_value="section").clean()
        with self.assertRaises(forms.ValidationError):
            FormField(type="statictext", default_value="p onclick=alert(1)").clean()

        quoted = FormField(type="statictext", title="T", choices='a"b').formfield()
        self.assertHTMLEqual(
            quoted.widget.render("quoted", None), '<div class="a&quot;b">T</div>'
        )

        form_instance = form.form_class()(data={"subject": "Hi"})
        self.assertTrue(form_instance.is_valid())
        self.assertFalse(form_instance.fields["para"].required)
        self.assertEqual(form_instance.fields["intro"].label, "")

        self.assertHTMLEqual(
            str(form_instance["intro"]),
            '<div class="lead highlight">Hello <b>world</b></div>',
        )
        self.assertHTMLEqual(str(form_instance["para"]), "<p>Paragraph</p>")
        self.assertIn(
            '<div class="lead highlight">Hello <b>world</b></div>',
            form_instance.as_p(),
        )

        submission = FormSubmission.objects.create(
            form=form, data=form_instance.cleaned_data
        )
        self.assertEqual(
            submission.data, {"subject": "Hi", "intro": None, "para": None}
        )

        # No-data fields are neither shown as current nor as removed fields
        self.assertEqual(
            [
                field["name"]
                for field in form.submissions_data(submissions=[submission])[0]["data"]
            ],
            ["subject"],
        )
        data = submission.formatted_data()
        self.assertEqual(data, "Subject:\nHi\n")
        self.assertNotIn("removed field", data)

        User.objects.create_superuser("admin", "admin@example.com", "password")
        self.client.login(username="admin", password="password")
        response = self.client.get(
            f"/admin/form_designer/form/{form.id}/export_submissions/"
        )
        self.assertEqual(response.status_code, 200)
        sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        self.assertEqual(rows[1][:1], ["subject"])
        self.assertEqual(len(rows[1]), 3)  # + submitted at, URL
